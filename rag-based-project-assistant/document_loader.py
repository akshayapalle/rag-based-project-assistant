
from pathlib import Path
from docx import Document

# Find the dataset folder beside this Python file
BASE_DIR = Path(__file__).resolve().parent
DATASET_DIR = BASE_DIR / "dataset"


def extract_docx_text(file_path):
    """Extract text from a Word document."""
    document = Document(file_path)
    parts = []

    # Read paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())

    # Read tables
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_txt_text(file_path):
    """Extract text from a plain-text file."""
    return file_path.read_text(encoding="utf-8")


def load_documents():
    """Load all .txt and .docx documents from the dataset."""
    documents = []

    for file_path in DATASET_DIR.rglob("*"):
        if not file_path.is_file():
            continue

        if file_path.suffix.lower() == ".docx":
            text = extract_docx_text(file_path)
        elif file_path.suffix.lower() == ".txt":
            text = extract_txt_text(file_path)
        else:
            continue

        if text.strip():
            documents.append({
                "document_id": str(file_path.relative_to(DATASET_DIR)),
                "file_name": file_path.name,
                "text": text
            })

    return documents


if __name__ == "__main__":
    documents = load_documents()

    print(f"\nDocuments loaded: {len(documents)}")

    for document in documents:
        print(f"\n--- {document['file_name']} ---")
        print(document["text"][:300])
        print("...")
