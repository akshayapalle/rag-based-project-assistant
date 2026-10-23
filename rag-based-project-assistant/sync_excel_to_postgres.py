
import os
from pathlib import Path

import psycopg
from dotenv import load_dotenv
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "dataset" / "project_data"

load_dotenv(BASE_DIR / ".env")


def read_excel(filename, sheet_name):
    path = DATA_DIR / filename

    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [
            str(value).strip().lower().replace(" ", "_")
            if value is not None else ""
            for value in next(rows)
        ]

        records = []
        for values in rows:
            if not any(value is not None for value in values):
                continue

            record = dict(zip(headers, values))
            records.append(record)

        return records
    finally:
        workbook.close()


def required(record, *columns):
    missing = [column for column in columns if column not in record]
    if missing:
        raise ValueError(f"Missing Excel columns: {missing}")


def get_value(record, column):
    value = record.get(column)
    return value.strip() if isinstance(value, str) else value


def sync_data():
    project_rows = read_excel("project.xlsx", "Project")
    team_rows = read_excel("team.xlsx", "Team")
    sprint_rows = read_excel("sprints.xlsx", "Sprints")
    task_rows = read_excel("tasks.xlsx", "Tasks")

    # Validate required columns before changing the database.
    for row in project_rows:
        required(row, "project_id", "project_name")

    for row in team_rows:
        required(row, "member_id", "name")

    for row in sprint_rows:
        required(row, "sprint_id", "project_id", "sprint_name")

    for row in task_rows:
        required(row, "task_id", "project_id", "task_name")

    connection_string = {
        "host": os.getenv("DB_HOST", "localhost"),
        "port": int(os.getenv("DB_PORT", "5432")),
        "dbname": os.getenv("DB_NAME"),
        "user": os.getenv("DB_USER"),
        "password": os.getenv("DB_PASSWORD"),
    }

    if not all(connection_string.values()):
        raise ValueError("Check DB_HOST, DB_PORT, DB_NAME, DB_USER and DB_PASSWORD in .env")

    with psycopg.connect(**connection_string) as conn:
        with conn.cursor() as cur:

            # 1. Projects
            for row in project_rows:
                cur.execute("""
                    INSERT INTO projects
                    (project_id, project_name, description, status,
                     start_date, target_end_date, client)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    ON CONFLICT (project_id) DO UPDATE SET
                        project_name = EXCLUDED.project_name,
                        description = EXCLUDED.description,
                        status = EXCLUDED.status,
                        start_date = EXCLUDED.start_date,
                        target_end_date = EXCLUDED.target_end_date,
                        client = EXCLUDED.client
                """, (
                    get_value(row, "project_id"),
                    get_value(row, "project_name"),
                    get_value(row, "description"),
                    get_value(row, "status"),
                    get_value(row, "start_date"),
                    get_value(row, "target_end_date"),
                    get_value(row, "client"),
                ))

            # 2. Team members
            for row in team_rows:
                cur.execute("""
                    INSERT INTO team_members
                    (member_id, name, role, team, email)
                    VALUES (%s, %s, %s, %s, %s)
                    ON CONFLICT (member_id) DO UPDATE SET
                        name = EXCLUDED.name,
                        role = EXCLUDED.role,
                        team = EXCLUDED.team,
                        email = EXCLUDED.email
                """, (
                    get_value(row, "member_id"),
                    get_value(row, "name"),
                    get_value(row, "role"),
                    get_value(row, "team"),
                    get_value(row, "email"),
                ))

            # Build a name-to-ID map for task assignees.
            cur.execute("SELECT member_id, name FROM team_members")
            members = {
                name.strip().casefold(): member_id
                for member_id, name in cur.fetchall()
            }

            # 3. Sprints
            for row in sprint_rows:
                cur.execute("""
                    INSERT INTO sprints
                    (sprint_id, project_id, sprint_name,
                     start_date, end_date, status)
                    VALUES (%s, %s, %s, %s, %s, %s)
                    ON CONFLICT (sprint_id) DO UPDATE SET
                        project_id = EXCLUDED.project_id,
                        sprint_name = EXCLUDED.sprint_name,
                        start_date = EXCLUDED.start_date,
                        end_date = EXCLUDED.end_date,
                        status = EXCLUDED.status
                """, (
                    get_value(row, "sprint_id"),
                    get_value(row, "project_id"),
                    get_value(row, "sprint_name"),
                    get_value(row, "start_date"),
                    get_value(row, "end_date"),
                    get_value(row, "status"),
                ))

            # 4. Tasks
            for row in task_rows:
                assignee = get_value(row, "assignee")

                if assignee is not None:
                    assignee = str(assignee).strip()
                    if assignee not in {
                        member_id for member_id in members.values()
                    }:
                        assignee = members.get(assignee.casefold())

                        if assignee is None:
                            raise ValueError(
                                f"Unknown task assignee: {row.get('assignee')}"
                            )

                cur.execute("""
                    INSERT INTO tasks
                    (task_id, project_id, task_name, description,
                     assignee, status, priority, deadline, sprint_id)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                    ON CONFLICT (task_id) DO UPDATE SET
                        project_id = EXCLUDED.project_id,
                        task_name = EXCLUDED.task_name,
                        description = EXCLUDED.description,
                        assignee = EXCLUDED.assignee,
                        status = EXCLUDED.status,
                        priority = EXCLUDED.priority,
                        deadline = EXCLUDED.deadline,
                        sprint_id = EXCLUDED.sprint_id
                """, (
                    get_value(row, "task_id"),
                    get_value(row, "project_id"),
                    get_value(row, "task_name"),
                    get_value(row, "description"),
                    assignee,
                    get_value(row, "status"),
                    get_value(row, "priority"),
                    get_value(row, "deadline"),
                    get_value(row, "sprint_id"),
                ))

    print("Excel-to-PostgreSQL sync completed successfully.")
    print("Projects, team members, sprints and tasks were synchronized.")
    print("Existing records missing from Excel were not deleted.")
    print("Change requests were not modified.")


if __name__ == "__main__":
    try:
        sync_data()
    except Exception as error:
        print(f"Sync failed: {error}")
